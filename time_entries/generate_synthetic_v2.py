"""
Synthetic time-entry generator v2 (corporate/transactional practice).
All clients, people, and matters are fictional. Seeded for reproducibility.

Design goals (fixes from v1 review):
- Narrative paraphrases: many variants per task, tasks reused across matters of the same type.
- Train/test split by narrative variant (held-out phrasings) + held-out client misspellings.
- Matter Master lists only legitimate aliases; misspellings are generated, never listed.
- Labels computed by one oracle, so Expected Action and Resolution Basis always agree.
- Explicit NEEDS_REVIEW / SPLIT_ENTRY / FLAG_MISMATCH actions instead of forced guesses.
- Hard cases: off-history staffing, multi-matter entries, closed matters, invalid CMIDs,
  same-client wrong matter, generic narratives that fit several matters.
- Hours drawn from task-specific skewed ranges (0.1 increments, long closing days possible).
"""
import random, re, datetime as dt
import pandas as pd

SEED = 210
rng = random.Random(SEED)

# ---------------------------------------------------------------- matters
M = [  # cmid, client, matter, type, status, codename, listed aliases, entities
 ("1111.001","Firm","Admin","ADMIN","Open","",["Firm"],[]),
 ("1111.002","Firm","Business Development","BD","Open","",["Firm"],[]),
 ("2201.002","Harborlight Bio","Seed Financing","VC","Closed","",["Harborlight","HLB"],["Seed"]),
 ("2201.003","Harborlight Bio","Series A Financing","VC","Open","",["Harborlight","HLB"],["Series A","Arcadia Ventures","Arcadia"]),
 ("2956.001","Solace Media","Licensing - Content Distribution","LICENSING","Open","",["Solace"],["Northwind Studios","Northwind"]),
 ("3012.001","Halvorsen Industrial","General Corporate","GENCORP","Open","",["Halvorsen","Halvorsen Ind."],[]),
 ("3012.004","Halvorsen Industrial","Acquisition of Brightline Pumps","MA_BUY","Open","Project Brightline",["Halvorsen","Halvorsen Ind."],["Brightline","Project Brightline"]),
 ("4471.001","Ridgewater Partners","General Fund Matters","FUND_GEN","Open","",["Ridgewater","RWP"],["Fund II"]),
 ("4471.002","Ridgewater Partners","Fund III Formation","FUND_FORM","Open","",["Ridgewater","RWP"],["Fund III","Calpine Pension"]),
 ("4590.001","Tidewater AI","General Corporate","GENCORP","Open","",["Tidewater"],[]),
 ("4590.002","Tidewater AI","Board Matters","BOARD","Open","",["Tidewater"],[]),
 ("4590.003","Tidewater AI","Series B Financing","VC","Open","Project Kite",["Tidewater"],["Series B","Summit Growth","Project Kite"]),
 ("5520.001","Cobalt Ridge Capital","Add-On Acquisition (Project Lark)","MA_BUY","Open","Project Lark",["Cobalt Ridge","CRC"],["Project Lark","Lark","Larkspur Controls"]),
 ("5520.002","Cobalt Ridge Capital","Portfolio Co. Sale (Project Wren)","MA_SELL","Open","Project Wren",["Cobalt Ridge","CRC"],["Project Wren","Wren","Wrenfield Logistics"]),
 ("6103.001","Meridian Therapeutics","SEC Reporting","SEC","Open","",["Meridian","Meridian Tx"],[]),
 ("6103.002","Meridian Therapeutics","Follow-On Offering","CAPMKT","Open","Project Ember",["Meridian","Meridian Tx"],["Project Ember","Hartwell"]),
 ("6103.003","Meridian Therapeutics","Board & Governance","BOARD","Open","",["Meridian","Meridian Tx"],[]),
 ("7215.002","Northgate Bank","Credit Facility - Pinecrest Mfg.","DEBT","Closed","",["Northgate","NGB"],["Pinecrest"]),
 ("7215.003","Northgate Bank","Credit Facility - Asterion Holdings","DEBT","Open","",["Northgate","NGB"],["Asterion"]),
 ("7215.004","Northgate Bank","Credit Facility - Kestrel Foods","DEBT","Open","",["Northgate","NGB"],["Kestrel"]),
 ("8340.001","Pelican Row Holdings","Acquisition of 450 Harbor St.","RE_ACQ","Open","",["Pelican Row","PRH"],["450 Harbor","Harbor St."]),
 ("8340.002","Pelican Row Holdings","Lease - 12 Mission Ave.","RE_LEASE","Open","",["Pelican Row","PRH"],["12 Mission","Mission Ave."]),
 ("9028.001","Lanternfish Software","General Corporate","GENCORP","Open","",["Lanternfish","LFS"],[]),
 ("9028.002","Lanternfish Software","Tax Structuring - Reorg","TAX","Open","Project Tern",["Lanternfish","LFS"],["Project Tern","Tern"]),
 ("9028.003","Lanternfish Software","Exec Comp / 409A","EXECCOMP","Open","",["Lanternfish","LFS"],[]),
]
MAT = {m[0]: dict(cmid=m[0],client=m[1],matter=m[2],type=m[3],status=m[4],code=m[5],aliases=m[6],ents=m[7]) for m in M}
OPEN = [c for c in MAT if MAT[c]["status"]=="Open"]
CLIENTS = sorted({m[1] for m in M if m[1]!="Firm"})

# ------------------------------------------------------------- timekeepers
TK = {  # id: (name, role, {cmid: 90-day hours})
 "TK101":("A. Chen","Partner",{"1111.001":6.0,"1111.002":14.5,"3012.001":7.2,"3012.004":62.4,"4590.002":9.5,"5520.001":48.1,"5520.002":21.7}),
 "TK102":("R. Patel","Sr. Associate",{"1111.001":5.1,"3012.001":8.2,"3012.004":118.6,"5520.001":74.3}),
 "TK103":("L. Moreno","Associate",{"1111.001":4.4,"4471.001":12.4,"4471.002":96.8,"5520.002":33.9}),
 "TK104":("D. Okafor","Associate",{"1111.001":3.8,"6103.001":41.2,"6103.002":88.0,"6103.003":16.3}),
 "TK105":("J. Whitfield","Associate",{"1111.001":4.0,"7215.003":71.5,"7215.004":19.2,"8340.001":55.0,"8340.002":7.6}),
 "TK106":("S. Kim","Associate",{"1111.001":6.3,"1111.002":3.5,"2201.003":46.9,"2956.001":18.4,"4590.001":30.2,"9028.001":22.3,"9028.002":14.1,"9028.003":10.8}),
 "TK107":("M. Alvarez","Paralegal",{"1111.001":9.0,"2201.003":28.0,"3012.004":37.1,"6103.001":19.5,"8340.001":24.2}),
 "TK108":("P. Hughes","Partner",{"1111.002":11.0,"2201.003":14.6,"4590.001":8.9,"4590.002":12.2,"4590.003":31.5,"6103.003":6.4,"9028.001":5.2}),
 "TK109":("E. Brandt","Associate",{"1111.001":5.5,"2201.003":22.4,"2956.001":15.0,"4471.002":27.3,"4590.003":58.8}),
 "TK110":("N. Farouk","Paralegal",{"1111.001":7.7,"4471.002":30.1,"4590.003":16.6,"5520.001":21.0,"7215.003":18.3,"7215.004":12.9}),
}

# ------------------------------------------------------------------ tasks
# types, quality, hours lo/hi, paralegal-ok, variants. Slots: {E} matter entity, {CP} counterparty side.
ALL_WORK = ["MA_BUY","MA_SELL","VC","FUND_FORM","FUND_GEN","CAPMKT","SEC","BOARD","DEBT","RE_ACQ","RE_LEASE","LICENSING","GENCORP","TAX","EXECCOMP"]
T = {
 "ma_spa":(["MA_BUY","MA_SELL"],"SPECIFIC",1.0,5.5,False,[
   "Review and mark up {E} purchase agreement","review and markup SPA; revise indemnification provisions",
   "Revise purchase agreement per client comments","markup SPA draft re indemnity and escrow",
   "Turn comments on {E} purchase agreement","revise definitive agreement and circulate to client",
   "Mark up APA for {E} deal","revise reps and warranties in purchase agreement"]),
 "ma_ds_draft":(["MA_SELL"],"SPECIFIC",1.0,5.0,False,[
   "Draft disclosure schedules for {E} sale","draft disclosure schedules; telephone conference with client re same",
   "update disclosure schedules per management input","revise disclosure schedules re material contracts",
   "prepare seller disclosure schedules","incorporate company comments into disclosure schedules"]),
 "ma_ds_rvw":(["MA_BUY"],"SPECIFIC",0.8,3.5,False,[
   "Review seller disclosure schedules and flag issues","review {E} disclosure schedules",
   "comment on seller disclosure schedules","prepare issues list re disclosure schedules",
   "cross-check disclosure schedules against diligence findings"]),
 "ma_dd":(["MA_BUY"],"SPECIFIC",0.6,6.5,True,[
   "Review {E} material contracts in data room; update diligence memo","Update due diligence tracker",
   "diligence review of {E} IP and employment agreements","draft diligence memo sections re commercial contracts",
   "review target change of control provisions","update diligence request list and follow up on open items",
   "summarize {E} leases for diligence memo"]),
 "ma_vdr":(["MA_SELL"],"SPECIFIC",0.5,4.0,True,[
   "Coordinate data room setup for {E} sale process","populate data room and index documents",
   "update data room index and upload responsive documents","respond to buyer diligence requests",
   "organize {E} data room folders"]),
 "ma_ppa":(["MA_BUY","MA_SELL"],"SPECIFIC",0.4,2.0,False,[
   "telephone conference with {CP} counsel re purchase price adjustment and working capital peg",
   "call with {CP} counsel re net working capital peg","conference with client re purchase price adjustment mechanics",
   "review {CP} comments on closing statement and working capital","draft email to client re escrow and holdback terms"]),
 "ma_loi":(["MA_BUY"],"SPECIFIC",0.5,2.5,False,[
   "Review {E} LOI and comment","comment on LOI terms re exclusivity","revise letter of intent per client call",
   "review draft term sheet for acquisition"]),
 "ma_cim":(["MA_SELL"],"SPECIFIC",0.8,3.5,False,[
   "Draft legal section of {E} CIM","review CIM legal disclosures","revise confidential information memorandum risk language"]),
 "ma_fin":(["MA_BUY"],"SPECIFIC",0.4,2.5,False,[
   "telephone conference with lender counsel re acquisition financing for {E}","review debt commitment letter for acquisition financing",
   "coordinate financing conditions with lender counsel"]),
 "closing":(["MA_BUY","MA_SELL","VC","DEBT","RE_ACQ","CAPMKT"],"SPECIFIC",0.5,4.5,True,[
   "Prepare signature packets for {E} closing","assemble signature packets","update closing checklist",
   "Prepare closing checklist for {E}","coordinate release of signature pages","compile executed closing documents",
   "prepare closing set index"]),
 "closing_day":(["MA_BUY","MA_SELL","VC","DEBT","RE_ACQ"],"SPECIFIC",4.5,10.5,False,[
   "Attend to {E} closing; coordinate funds flow and release of signatures",
   "manage closing: finalize ancillary documents, confirm conditions satisfied, release signatures",
   "closing day for {E}: resolve open items with opposing counsel, circulate final documents"]),
 "vc_docs":(["VC"],"SPECIFIC",1.0,5.0,False,[
   "revise SPA and IRA per lead investor comments","draft voting agreement and ROFR agreement",
   "mark up {E} term sheet","revise NVCA documents per {E} comments","prepare charter amendment for {E}",
   "revise investors' rights agreement re information rights"]),
 "vc_cap":(["VC"],"SPECIFIC",0.5,3.0,True,[
   "Review cap table and pro forma; flag option pool issue","update pro forma cap table for {E}",
   "reconcile cap table to stock ledger","prepare {E} capitalization table"]),
 "vc_side":(["VC"],"SPECIFIC",0.3,1.8,False,[
   "telephone conference with {E} counsel re lead investor side letter","draft management rights letter",
   "revise side letter per investor comments"]),
 "fund_lpa":(["FUND_FORM"],"SPECIFIC",0.8,5.0,False,[
   "Review LPA comments and circulate to team","revise LPA waterfall provisions","mark up {E} LPA",
   "revise LPA key person and removal provisions","prepare LPA comment chart"]),
 "fund_side":(["FUND_FORM"],"SPECIFIC",0.5,3.5,False,[
   "Draft side letter for LP ({E}); conference re MFN elections","prepare MFN election forms",
   "review LP side letter requests","compile MFN compendium for {E}"]),
 "fund_sub":(["FUND_FORM"],"SPECIFIC",0.5,4.0,True,[
   "Prepare subscription documents and investor questionnaire","process LP subscription documents",
   "review subscription documents for completeness","track {E} investor subscriptions"]),
 "fund_ppm":(["FUND_FORM"],"SPECIFIC",0.8,4.0,False,[
   "revise PPM risk factors","update {E} PPM","draft PPM conflicts of interest section"]),
 "fund_gen":(["FUND_GEN"],"SPECIFIC",0.3,2.5,True,[
   "review {E} capital call notice","prepare capital call notices","respond to LP inquiry re {E} distributions",
   "draft LPAC consent","update investor register"]),
 "cm_ua":(["CAPMKT"],"SPECIFIC",0.5,3.0,False,[
   "Review underwriting agreement draft","comment on underwriting agreement from {E} counsel",
   "revise underwriting agreement indemnification section"]),
 "cm_lockup":(["CAPMKT"],"SPECIFIC",0.2,1.5,True,[
   "telephone conference with underwriters' counsel re lock-up agreements","track lock-up agreement signatures",
   "circulate lock-up agreements to directors and officers"]),
 "cm_s3":(["CAPMKT"],"SPECIFIC",1.0,5.0,False,[
   "Review S-3 draft and comment on risk factors","review comfort letter bring-down",
   "draft prospectus supplement sections","revise {E} prospectus supplement"]),
 "cm_dd":(["CAPMKT"],"SPECIFIC",0.5,2.0,False,[
   "due diligence call with management and underwriters","prepare for bring-down diligence call",
   "prepare {E} diligence call agenda"]),
 "sec_10q":(["SEC"],"SPECIFIC",0.8,4.5,True,[
   "Review 10-Q draft MD&A","comment on 10-Q","prepare 10-Q exhibits","update 10-Q subsequent events disclosure"]),
 "sec_8k":(["SEC"],"SPECIFIC",0.4,2.0,False,[
   "revise 8-K re officer departure; circulate to CFO","draft 8-K re material agreement",
   "prepare Item 5.02 disclosure"]),
 "sec_16":(["SEC"],"SPECIFIC",0.2,1.2,True,[
   "prepare Section 16 filings (Form 4s)","file Form 4 for director grant","update Section 16 tracking chart"]),
 "bd_mtg":(["BOARD"],"SPECIFIC",0.5,3.0,False,[
   "Attend board meeting","draft minutes of board meeting","prepare board meeting agenda and materials",
   "draft unanimous written consent of the board","prepare board deck legal update"]),
 "bd_gov":(["BOARD"],"SPECIFIC",0.5,2.5,False,[
   "Review committee charters and flag updates","revise corporate governance guidelines",
   "prepare D&O questionnaire"]),
 "debt":(["DEBT"],"SPECIFIC",0.6,5.0,True,[
   "Revise {E} credit agreement per borrower comments","mark up credit agreement","review {E} perfection certificate",
   "revise closing checklist; telephone conference with bank re conditions precedent","draft commitment letter for {E} facility",
   "review security agreement and UCC filings","prepare UCC search requests for {E}"]),
 "re_acq":(["RE_ACQ"],"SPECIFIC",0.5,4.5,True,[
   "review title commitment and survey; prepare title objection letter","review {E} title commitment",
   "review tenant estoppels","review {E} PSA amendment","prepare closing binder","review zoning report for {E}"]),
 "re_lease":(["RE_LEASE"],"SPECIFIC",0.5,3.5,False,[
   "mark up {E} lease; revise SNDA","review landlord comments to lease","draft lease amendment re TI allowance",
   "negotiate {E} lease renewal terms"]),
 "lic":(["LICENSING"],"SPECIFIC",0.5,4.0,False,[
   "mark up license agreement; revise IP ownership and termination provisions","review {E} comments to content license",
   "draft distribution agreement exhibits","revise territory and exclusivity provisions in {E} license"]),
 "gc":(["GENCORP"],"SPECIFIC",0.2,2.5,False,[
   "telephone conference with CEO re hiring questions","telephone conference with founder re equity questions",
   "draft consulting agreement","advise re option grant mechanics","review NDA for client",
   "prepare annual stockholder consent"]),
 "tax":(["TAX"],"SPECIFIC",0.5,3.5,False,[
   "review reorganization step plan","telephone conference with tax re {E} structure",
   "draft contribution agreement for reorganization","revise {E} step plan per tax comments"]),
 "ec":(["EXECCOMP"],"SPECIFIC",0.4,3.0,False,[
   "revise executive offer letter and severance terms","review 409A valuation report","draft RSU award agreement",
   "revise change in control severance plan"]),
 "admin":(["ADMIN"],"SPECIFIC",0.3,2.5,True,[
   "attend to administrative matters","update matter list and time entries for week","attend CLE training",
   "attend firm meeting"]),
 "bizdev":(["BD"],"SPECIFIC",0.5,4.0,False,[
   "Draft pitch for RFP; conference with marketing re same","prepare client alert","prepare for business development meeting"]),
 # vague: typed hint (narrows type) vs generic (no signal)
 "v_typed":(None,"VAGUE",0.2,3.0,True,None),
 "v_generic":(ALL_WORK,"VAGUE",0.1,2.5,True,[
   "attend to matter","review documents","emails re same","work on matter","attend to various issues",
   "review correspondence","follow up on open items"]),
}
V_TYPED = {  # type-hinted vague narratives
 "MA_BUY":["attend to deal matters","attend to diligence"], "MA_SELL":["attend to deal matters","attend to sale process"],
 "VC":["attend to financing matters"], "FUND_FORM":["attend to fund formation matters"], "FUND_GEN":["attend to fund matters"],
 "CAPMKT":["attend to offering matters"], "SEC":["attend to SEC matters"], "BOARD":["attend to board matters"],
 "DEBT":["attend to loan documents"], "RE_ACQ":["attend to property acquisition"], "RE_LEASE":["attend to lease matters"],
 "LICENSING":["attend to license matters"], "GENCORP":["attend to corporate matters"], "TAX":["attend to tax matters"],
 "EXECCOMP":["attend to comp matters"],
}
V_TYPED_TYPES = {}
for ty, lst in V_TYPED.items():
    for s in lst: V_TYPED_TYPES.setdefault(s, set()).add(ty)

# ------------------------------------------------------ split by variant
HELDOUT = {}  # (task, variant_idx) -> "TEST"/"TRAIN"
for tid, t in T.items():
    vs = t[5]
    if not vs: continue
    n_hold = max(1, round(len(vs)*0.3))
    idx = list(range(len(vs))); rng.shuffle(idx)
    for i in idx[:n_hold]: HELDOUT[(tid,i)] = "TEST"
for ty, lst in V_TYPED.items():
    for i, s in enumerate(lst):
        HELDOUT[("v_typed", s)] = "TEST" if (len(lst)>1 and i==len(lst)-1) else "TRAIN"

# ------------------------------------------------------------ text noise
ABBR = [("telephone conference with","tcw"),("telephone conference with","t/c w/"),("conference with","conf w/"),
        ("review","rvw"),("review","revw"),("agreement","agmt"),("agreement","agrmt"),("regarding","re"),
        ("documents","docs"),("disclosure schedules","discl. schedules"),("signature packets","sig pckts"),
        ("diligence","DD"),("due diligence","DD"),("with","w/"),(" and "," & "),(" and "," + "),
        ("prepare","prep"),("comments","cmts"),("management","mgmt"),("subscription","sub"),
        ("purchase price adjustment","PPA"),("working capital","WC"),("board meeting","BOD mtg"),
        ("letter of intent","LOI"),("information","info"),("provisions","provs"),("draft","drft")]
def typo(w):
    if len(w) < 5 or not w.isalpha(): return w
    i = rng.randrange(1, len(w)-2); op = rng.random()
    if op < .4: return w[:i]+w[i+1]+w[i]+w[i+2:]
    if op < .7: return w[:i]+w[i+1:]
    return w[:i]+w[i]+w[i:]
def noise(s):
    for full, ab in ABBR:
        if full in s and rng.random() < 0.45: s = s.replace(full, ab, 1)
    words = s.split(" ")
    words = [typo(w) if rng.random() < 0.05 else w for w in words]
    s = " ".join(words)
    r = rng.random()
    if r < 0.35: s = s[0].lower()+s[1:]
    elif r < 0.45: s = s.upper() if len(s) < 25 else s
    if rng.random() < 0.15: s += "."
    return s

# ---------------------------------------------------- client surface forms
ABBRW = {"Industrial":["Ind","Indus.","Industrl"],"Therapeutics":["Thera","Therap.","Therapeutic"],"Software":["Sftwr","Soft."],
         "Holdings":["Hldgs","Holding"],"Partners":["Ptnrs","Partner"],"Capital":["Cap.","Cap"],"Bank":["Bnk","Bk"],
         "Media":["Med.","Media Inc"],"Bio":["Biosciences","Biotech"],"AI":["A.I.","ai",".ai"]}
def client_variants(c):
    listed = set([c] + [a for m in MAT.values() if m["client"]==c for a in m["aliases"]])
    pool = set(); ws = c.split(); tries = 0
    while len(pool) < 10 and tries < 500:
        tries += 1; op = rng.random(); w = list(ws)
        if op < .3:
            j = rng.randrange(len(w)); w[j] = typo(w[j]) if len(w[j])>=5 else w[j]; v = " ".join(w)
        elif op < .55 and w[-1] in ABBRW:
            w[-1] = rng.choice(ABBRW[w[-1]]); v = " ".join(w).replace(" .ai",".ai")
        elif op < .7: v = typo(w[0])
        elif op < .85: v = " ".join(w).lower()
        else: v = w[0][:max(4,len(w[0])-3)]
        if v not in listed and v.strip(): pool.add(v)
    pool = sorted(pool); rng.shuffle(pool)
    h = len(pool)//2
    return pool[:h], pool[h:]
CV = {c: client_variants(c) for c in CLIENTS}
def client_surface(c, split):
    if c == "Firm": return "Firm", "CANONICAL"
    r = rng.random()
    a, b = (.15, .35) if split == "TEST" else (.25, .55)  # test leans on unseen misspellings
    if r < a: return c, "CANONICAL"
    if r < b: return rng.choice([a for a in MAT_BY_CLIENT[c][0]["aliases"]]), "LISTED_ALIAS"
    pool = CV[c][1] if split=="TEST" else CV[c][0]
    return rng.choice(pool), "UNLISTED_VARIANT"
MAT_BY_CLIENT = {}
for m in MAT.values(): MAT_BY_CLIENT.setdefault(m["client"], []).append(m)

# ------------------------------------------------------------- CMID forms
def cmid_bad_format(c):
    a, b = c.split("."); n = int(b)
    opts = [f"{a}-{b}", f"{a}{b}", f"{a}.{n}", f"{a}.{n:02d}", f"{a} .{b}", f"{a}/{b}"]
    return rng.choice(opts)
def normalize_cmid(raw):
    if not raw: return None
    d = re.sub(r"[^0-9.\-/ ]","",raw).strip()
    m = re.match(r"^(\d{4})\s*[.\-/]\s*(\d{1,3})$", d)
    if m: c = f"{m.group(1)}.{int(m.group(2)):03d}"
    elif re.fullmatch(r"\d{7}", d): c = f"{d[:4]}.{d[4:]}"
    else: return None
    return c if c in MAT else None

# ------------------------------------------------------ narrative builder
def pick_task(mtype, role, allow_vague=True):
    opts = [tid for tid,t in T.items() if t[0] and mtype in t[0] and t[1]=="SPECIFIC" and tid!="closing_day" and (role!="Paralegal" or t[4])]
    return rng.choice(opts)
def render(tid, cm, want_split=None):
    """returns text, used_entity(bool), split, template_id"""
    m = MAT[cm]; t = T[tid]
    if tid == "v_typed":
        s = rng.choice(V_TYPED[m["type"]]); return noise(s), False, HELDOUT[("v_typed",s)], f"v_typed:{s}"
    vs = t[5]
    cand = [i for i,v in enumerate(vs) if ("{E}" not in v or m["ents"])]
    if want_split:
        c2 = [i for i in cand if HELDOUT.get((tid,i),"TRAIN")==want_split]
        cand = c2 or cand
    i = rng.choice(cand); v = vs[i]
    used = "{E}" in v
    v = noise(v)
    if used: v = v.replace("{E}", rng.choice(m["ents"]))
    side = "seller" if m["type"]=="MA_BUY" else "buyer"
    v = v.replace("{CP}", side)
    tp = f"{tid}:{i}" + (":CP" if "{CP}" in vs[i] else "")
    return v, used, HELDOUT.get((tid,i),"TRAIN"), tp
def types_of(tid, text_hint=None, mtype=None):
    if tid == "v_typed": return set(V_TYPED_TYPES[text_hint])
    if text_hint and text_hint.endswith(":CP") and mtype: return {mtype}  # 'seller/buyer counsel' reveals deal side
    return set(T[tid][0])

# ------------------------------------------------------------------ oracle
def oracle(e):
    """Decides expected action from what a reviewer could know: entered fields, narrative entity
    mentions, task-type compatibility, Matter Master, TK history. Never uses hours to break ties."""
    tk = e["tk"]; hist = [c for c in TK[tk][2] if MAT[c]["status"]=="Open"]
    ent = e["ent_matters"]; types = e["types"]
    if e["split_matters"]:
        return "SPLIT_ENTRY", "Narrative names work on two matters ("+" & ".join(e["split_matters"])+"); split entry"
    raw = e["cmid_raw"]; norm = normalize_cmid(raw)
    if raw and norm:
        mm = MAT[norm]
        if mm["status"] == "Closed":
            basis = f"Entered CMID {norm} is a closed matter"
        elif (ent and norm not in ent) or (mm["type"] not in types):
            basis = f"Entered CMID {norm} ({mm['matter']}) contradicts narrative"
        else:
            b = "Entered CMID" if raw == norm else f"Normalize CMID format ({raw} -> {norm})"
            return "AUTO_MATCH", b
        # re-resolve from other signals
        rest, why = resolve_rest(e, hist, ent, types)
        if len(rest)==1: return "FLAG_MISMATCH", basis+f"; narrative supports {rest[0]} ({why})"
        return "FLAG_MISMATCH", basis+"; correct matter not determinable ("+", ".join(rest)+")"
    pre = f"Invalid CMID '{raw}' ignored; " if raw else ""
    rest, why = resolve_rest(e, hist, ent, types)
    if len(rest)==1: return "AUTO_MATCH", pre+why
    return "NEEDS_REVIEW", pre+f"AMBIGUOUS - {len(rest)} candidates ("+", ".join(rest)+"); "+why
def resolve_rest(e, hist, ent, types):
    why = []; cands = list(hist)
    if e["matter_raw"]:
        cm = [c for c in OPEN if MAT[c]["code"]==e["matter_raw"]]
        if cm: return cm, f"Codename '{e['matter_raw']}' maps to {cm[0]}"
    if e["client_canon"]:
        cl = [c for c in cands if MAT[c]["client"]==e["client_canon"]]
        if cl: cands = cl; why.append(f"client '{e['client_raw']}' -> {e['client_canon']}")
        else:
            cands = [c for c in OPEN if MAT[c]["client"]==e["client_canon"]]
            why.append(f"client '{e['client_raw']}' -> {e['client_canon']} (not in TK history)")
    if ent:
        c2 = [c for c in cands if c in ent]
        if c2: cands = c2; why.append("narrative names matter-specific term")
        else: cands = sorted(ent); why.append("narrative names matter-specific term; matter not in TK history")
    c3 = [c for c in cands if MAT[c]["type"] in types]
    if c3 and len(c3) < len(cands): why.append("task type narrows candidates")
    if c3: cands = c3
    if len(cands)==1 and not why: why.append("only compatible matter in TK history")
    if len(cands)==1: why.append(f"=> {cands[0]}")
    return sorted(cands), "; ".join(why) if why else "no client/matter signal"

# ------------------------------------------------------------------- dates
def biz_days(a, b):
    d = a; out = []
    while d <= b:
        if d.weekday() < 5: out.append(d)
        d += dt.timedelta(days=1)
    return out
START, END = dt.date(2026,8,3), dt.date(2026,9,18)
BD = biz_days(START, END)
WEEKEND = [d for d in (START+dt.timedelta(n) for n in range((END-START).days+1)) if d.weekday()>=5]

def hours(tid, role):
    lo, hi = T[tid][2], T[tid][3]
    x = lo + (hi-lo)*rng.betavariate(1.6, 3.2)
    if role == "Partner": x *= 0.8
    return max(0.1, round(x, 1))

# ------------------------------------------------------------ generation
INPUT_MIX = [("FULL",.34),("CMID_ONLY",.10),("CLIENT_ONLY",.20),("CODENAME",.05),("NONE",.18),("WRONG_CMID",.08),("INVALID_CMID",.03)]
def draw_input():
    r = rng.random(); a = 0
    for k,p in INPUT_MIX:
        a += p
        if r < a: return k
    return "NONE"

rows = []
def make_entry(tk, cm, forced=None, off_history=False):
    name, role, hist = TK[tk]; m = MAT[cm]
    want_split = "TEST" if rng.random() < 0.28 else "TRAIN"
    q = rng.random()
    # narrative quality
    if cm.startswith("1111"): quality = "SPECIFIC"
    elif q < 0.16: quality = "VAGUE"
    elif q < 0.30 and role != "Paralegal": quality = "BLOCK_BILLED"
    else: quality = "SPECIFIC"
    ent_used = False; tids = []; splits = []; tmpl = []
    if quality == "VAGUE":
        tid = "v_typed" if rng.random() < .55 else "v_generic"
        txt, ent_used, sp, tp = render(tid, cm, want_split); tids=[tid]; splits=[sp]; tmpl=[tp]
        types = types_of(tid, tp.split(":",1)[1]) if tid=="v_typed" else types_of(tid)
        hrs = hours(tid, role)
    elif quality == "BLOCK_BILLED":
        k = rng.choice([2,2,3]); parts = []; types = None; hrs = 0
        used_t = set()
        for _ in range(k):
            tid = pick_task(m["type"], role)
            if tid in used_t and len(used_t) < 3: tid = pick_task(m["type"], role)
            used_t.add(tid)
            t, u, sp, tp = render(tid, cm, want_split)
            if tp.split(":CP")[0] in [x.split(":CP")[0] for x in tmpl]: continue
            parts.append(t.rstrip(".")); ent_used |= u
            tids.append(tid); splits.append(sp); tmpl.append(tp)
            ty = types_of(tid, tp, m["type"]); types = ty if types is None else (types & ty or types)
            hrs += hours(tid, role)
        if len(parts) < 2: return None
        txt = "; ".join(p[0].lower()+p[1:] if (i and not p.isupper()) else p for i,p in enumerate(parts))
        hrs = round(min(hrs*0.95, 11.5), 1)
    else:
        tid = pick_task(m["type"], role)
        if rng.random() < 0.06 and m["type"] in ["MA_BUY","MA_SELL","VC","DEBT","RE_ACQ"] and role!="Paralegal": tid = "closing_day"
        txt, ent_used, sp, tp = render(tid, cm, want_split); tids=[tid]; splits=[sp]; tmpl=[tp]
        types = types_of(tid, tp, m["type"]); hrs = hours(tid, role)
    split = "TEST" if "TEST" in splits else "TRAIN"
    it = forced or draw_input()
    if it == "CODENAME" and not m["code"]: it = "NONE"
    if it == "NONE" and m["code"] and rng.random() < .35: it = "CODENAME"
    if off_history and it in ("FULL","CMID_ONLY","WRONG_CMID","INVALID_CMID"): it = rng.choice(["CLIENT_ONLY","NONE"])
    if off_history and it == "NONE" and not ent_used: return None
    e = dict(tk=tk, cm=cm, types=types, ent_matters={cm} if ent_used else set(), split_matters=None,
             cmid_raw="", client_raw="", client_canon="", matter_raw="", surface="", hard=[])
    if off_history: e["hard"].append("OFF_HISTORY")
    client_disp = ""; matter_disp = ""
    if it == "FULL":
        e["cmid_raw"] = cm; client_disp = m["client"]; matter_disp = m["matter"]; e["surface"]="CANONICAL"
    elif it == "CMID_ONLY":
        e["cmid_raw"] = cmid_bad_format(cm) if rng.random() < .5 else cm
    elif it == "CLIENT_ONLY":
        s, surf = client_surface(m["client"], split); e["client_raw"]=s; e["client_canon"]=m["client"]; e["surface"]=surf; client_disp=s
    elif it == "CODENAME":
        e["matter_raw"] = m["code"]; matter_disp = m["code"]
    elif it == "WRONG_CMID":
        r = rng.random()
        same_client = [c for c in MAT if MAT[c]["client"]==m["client"] and c!=cm and not c.startswith("1111")]
        if r < .4 and same_client:
            wrong = rng.choice(same_client); e["hard"].append("SAME_CLIENT_WRONG_MATTER" if MAT[wrong]["status"]=="Open" else "CLOSED_MATTER")
        else:
            others = [c for c in hist if c != cm and not c.startswith("1111")] or [c for c in OPEN if c!=cm]
            wrong = rng.choice(others)
            if MAT[wrong]["status"]=="Closed": e["hard"].append("CLOSED_MATTER")
        e["cmid_raw"] = wrong; client_disp = MAT[wrong]["client"]; matter_disp = MAT[wrong]["matter"]
    elif it == "INVALID_CMID":
        a, b = cm.split(".")
        bad = a[1]+a[0]+a[2:]+"."+b if rng.random()<.5 else a+"."+b[::-1]
        if normalize_cmid(bad): return None
        e["cmid_raw"] = bad; e["hard"].append("INVALID_CMID")
    e["matter_raw"] = e["matter_raw"]
    # CLOSED_MATTER specific: Harborlight Seed / Pinecrest typed out of habit
    action, basis = oracle(e)
    if it == "WRONG_CMID" and action != "FLAG_MISMATCH": return None  # undetectable wrong CMID: skip
    if action == "AUTO_MATCH" and "->" not in basis and it not in ("FULL","CMID_ONLY") and not basis:
        pass
    # consistency guard
    if action in ("AUTO_MATCH",) :
        rest = [cm]  # oracle must have landed on truth
        chk = normalize_cmid(e["cmid_raw"]) if e["cmid_raw"] else None
        if chk and chk != cm: return None
        if not chk:
            r2, _ = resolve_rest(e, [c for c in TK[tk][2] if MAT[c]["status"]=="Open"], e["ent_matters"], e["types"])
            if r2 != [cm]: return None
    if action == "FLAG_MISMATCH" and "narrative supports" in basis and f"supports {cm}" not in basis: return None
    if action == "NEEDS_REVIEW":
        r2, _ = resolve_rest(e, [c for c in TK[tk][2] if MAT[c]["status"]=="Open"], e["ent_matters"], e["types"])
        if cm not in r2: return None
        e["hard"].append("GENUINELY_AMBIGUOUS")
    if ("SPLIT" not in it) and rng.random() < 0.035 and m["type"] in ["MA_BUY","MA_SELL","DEBT","RE_ACQ","CAPMKT","VC"]:
        d = rng.choice(WEEKEND)
    else: d = rng.choice(BD)
    return dict(tk=tk, name=name, date=d, cmid=e["cmid_raw"], client=client_disp, matter=matter_disp, hours=hrs,
                narr=txt, true=cm, action=action, input=it, hard=";".join(e["hard"]), surface=e["surface"] or ("N/A"),
                quality=quality, tids=",".join(tids), tmpl="|".join(tmpl), split=split, basis=basis)

def make_split(tk):
    name, role, hist = TK[tk]
    cands = [c for c in hist if MAT[c]["ents"] and not c.startswith("1111")]
    if len(cands) < 2: return None
    a, b = rng.sample(cands, 2)
    parts = []; tids = []; tmpl = []; hrs = 0; splits=[]
    for cm in (a, b):
        tid = pick_task(MAT[cm]["type"], role)
        vs = T[tid][5]; ii = [i for i,v in enumerate(vs) if "{E}" in v]
        if not ii: return None
        i = rng.choice(ii); v = noise(vs[i]).replace("{E}", rng.choice(MAT[cm]["ents"])).replace("{CP}", "seller" if MAT[cm]["type"]=="MA_BUY" else "buyer")
        parts.append(v); tids.append(tid); tmpl.append(f"{tid}:{i}"); hrs += hours(tid, role); splits.append(HELDOUT.get((tid,i),"TRAIN"))
    txt = parts[0] + rng.choice(["; also ","; ",". Separately, "]) + parts[1][0].lower()+parts[1][1:]
    entered = rng.choice(["", a])
    e = dict(split_matters=[a,b], tk=tk)
    return dict(tk=tk, name=name, date=rng.choice(BD), cmid=entered, client=MAT[a]["client"] if entered else "", matter=MAT[a]["matter"] if entered else "",
                hours=round(hrs,1), narr=txt, true=f"{a} | {b}", action="SPLIT_ENTRY", input="FULL" if entered else "NONE",
                hard="MULTI_MATTER", surface="CANONICAL" if entered else "N/A", quality="BLOCK_BILLED", tids=",".join(tids),
                tmpl="|".join(tmpl), split="TEST" if "TEST" in splits else "TRAIN",
                basis=f"Narrative names work on two matters ({a} & {b}); split entry")

TARGET_PER_MATTER = 22
staff = {c: [tk for tk in TK if c in TK[tk][2]] for c in OPEN}
for cm in OPEN:
    n = 0; guard = 0
    while n < TARGET_PER_MATTER and guard < 2000:
        guard += 1
        tks = staff[cm]; w = [TK[t][2][cm] for t in tks]
        tk = rng.choices(tks, weights=w)[0]
        r = make_entry(tk, cm)
        if r: rows.append(r); n += 1
# extra volume weighted by history hours
pairs = [(tk, c, h) for tk in TK for c, h in TK[tk][2].items() if MAT[c]["status"]=="Open"]
while len(rows) < 600:
    tk, cm, _ = rng.choices(pairs, weights=[p[2] for p in pairs])[0]
    r = make_entry(tk, cm)
    if r: rows.append(r)
# off-history staffing
oh = 0
while oh < 16:
    tk = rng.choice(list(TK)); cm = rng.choice([c for c in OPEN if c not in TK[tk][2] and not c.startswith("1111")])
    if TK[tk][1]=="Paralegal" and not any(T[t][4] for t in T if T[t][0] and MAT[cm]["type"] in T[t][0] and t!="v_generic"): continue
    r = make_entry(tk, cm, off_history=True)
    if r and r["action"] in ("AUTO_MATCH","NEEDS_REVIEW"): rows.append(r); oh += 1
# closed-matter habit entries
for tk, closed, cur in [("TK106","2201.002","2201.003"),("TK108","2201.002","2201.003"),("TK105","7215.002","7215.003"),("TK110","7215.002","7215.004")]*2:
    for _ in range(50):
        r = make_entry(tk, cur, forced="FULL")
        if not r: continue
        e = dict(tk=tk, cm=cur, types=set(T[r["tids"].split(",")[0]][0] or []), ent_matters={cur} if any(x in r["narr"] for x in MAT[cur]["ents"]) else set(),
                 split_matters=None, cmid_raw=closed, client_raw="", client_canon="", matter_raw="", surface="", hard=[])
        a, b = oracle(e)
        if a == "FLAG_MISMATCH" and f"supports {cur}" in b:
            r.update(cmid=closed, client=MAT[closed]["client"], matter=MAT[closed]["matter"], action=a, basis=b,
                     input="WRONG_CMID", hard="CLOSED_MATTER"); rows.append(r); break
# multi-matter
sm = 0
while sm < 16:
    r = make_split(rng.choice(list(TK)))
    if r: rows.append(r); sm += 1

rows.sort(key=lambda r: (r["date"], r["tk"]))
out = []
for i, r in enumerate(rows):
    tc = r["true"].split(" | ")
    out.append({"Time ID": f"T-{2001+i:06d}","Status":"WIP","TK ID":r["tk"],"Timekeeper":r["name"],
        "Date":r["date"].strftime("%m/%d/%Y"),"CMID":r["cmid"],"Client":r["client"],"Matter":r["matter"],
        "Hours":r["hours"],"Narrative":r["narr"],
        "KEY: True CMID":r["true"],"KEY: True Client":MAT[tc[0]]["client"],
        "KEY: True Matter":" | ".join(MAT[c]["matter"] for c in tc),
        "KEY: Expected Action":r["action"],"KEY: Input Type":r["input"],"KEY: Hard Case":r["hard"],
        "KEY: Client Surface":r["surface"],"KEY: Narrative Quality":r["quality"],"KEY: Task IDs":r["tids"],
        "KEY: Template IDs":r["tmpl"],"KEY: Split":r["split"],"KEY: Resolution Basis":r["basis"]})
TE = pd.DataFrame(out)

MM = pd.DataFrame([{"CMID":m["cmid"],"Client":m["client"],"Matter":m["matter"],"Matter Type":m["type"],
    "Status":m["status"],"Codename":m["code"],"Known Aliases":", ".join(m["aliases"]),
    "Key Terms":", ".join(x for x in m["ents"] if x!=m["code"])} for m in MAT.values()])
HI = pd.DataFrame([{"TK ID":tk,"Timekeeper":v[0],"Role":v[1],"CMID":c,"Client":MAT[c]["client"],"Matter":MAT[c]["matter"],
    "Hours (last 90 days)":h} for tk,v in TK.items() for c,h in v[2].items()])

# ------------------------------------------------------------------ export
if __name__ == "__main__":
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    OUT = "synthetic_time_entries_v2.xlsx"
    wb = Workbook(); F = "Arial"
    blue = PatternFill("solid", fgColor="1F4E79"); gold = PatternFill("solid", fgColor="BF8F00"); grey = PatternFill("solid", fgColor="595959")
    def sheet(ws, df, key_prefix=None, widths=None):
        ws.append(list(df.columns))
        for r in df.itertuples(index=False): ws.append(["" if (isinstance(v,float) and pd.isna(v)) else v for v in r])
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(1, j); cell.font = Font(name=F, bold=True, color="FFFFFF")
            cell.fill = gold if (key_prefix and str(c).startswith(key_prefix)) else (blue if key_prefix else grey)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, max(10, min(45, int(df[c].astype(str).str.len().quantile(.9))+2)))
        for row in ws.iter_rows(min_row=2):
            for cell in row: cell.font = Font(name=F, size=10)
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    ws = wb.active; ws.title = "Time Entries"
    sheet(ws, TE, "KEY:", {"Narrative":60,"KEY: Resolution Basis":60,"KEY: True Matter":34})
    for c in ws["I"][1:]: c.number_format = "0.0"
    sheet(wb.create_sheet("Matter Master"), MM)
    sheet(wb.create_sheet("TK Matter History"), HI)

    s = wb.create_sheet("Summary")
    s["A1"] = "Dataset summary (live formulas over Time Entries)"; s["A1"].font = Font(name=F, bold=True, size=12)
    n = len(TE)+1; rngs = lambda col: f"'Time Entries'!${col}$2:${col}${n}"
    colmap = {c: get_column_letter(i+1) for i, c in enumerate(TE.columns)}
    r = 3
    blocks = [("Expected Action","KEY: Expected Action",["AUTO_MATCH","FLAG_MISMATCH","NEEDS_REVIEW","SPLIT_ENTRY"]),
              ("Input Type","KEY: Input Type",["FULL","CMID_ONLY","CLIENT_ONLY","CODENAME","NONE","WRONG_CMID","INVALID_CMID"]),
              ("Narrative Quality","KEY: Narrative Quality",["SPECIFIC","VAGUE","BLOCK_BILLED"]),
              ("Client Surface","KEY: Client Surface",["CANONICAL","LISTED_ALIAS","UNLISTED_VARIANT","N/A"])]
    sp = colmap["KEY: Split"]
    for title, field, vals in blocks:
        for j, h in enumerate([title, "TRAIN", "TEST", "Total"]):
            c = s.cell(r, 1+j, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = grey
        r += 1; top = r
        for v in vals:
            s.cell(r, 1, v).font = Font(name=F)
            s.cell(r, 2, f'=COUNTIFS({rngs(colmap[field])},$A{r},{rngs(sp)},"TRAIN")')
            s.cell(r, 3, f'=COUNTIFS({rngs(colmap[field])},$A{r},{rngs(sp)},"TEST")')
            s.cell(r, 4, f"=B{r}+C{r}"); r += 1
        s.cell(r, 1, "Total").font = Font(name=F, bold=True)
        for j, L in enumerate("BCD", 2): s.cell(r, j, f"=SUM({L}{top}:{L}{r-1})").font = Font(name=F, bold=True)
        r += 2
    s.cell(r, 1, "Hard cases (count of rows whose Hard Case contains)").font = Font(name=F, bold=True); r += 1
    for hc in ["GENUINELY_AMBIGUOUS","OFF_HISTORY","MULTI_MATTER","SAME_CLIENT_WRONG_MATTER","CLOSED_MATTER","INVALID_CMID"]:
        s.cell(r, 1, hc).font = Font(name=F); s.cell(r, 4, f'=COUNTIF({rngs(colmap["KEY: Hard Case"])},"*"&A{r}&"*")'); r += 1
    r += 1; s.cell(r, 1, "Hours").font = Font(name=F, bold=True); r += 1
    for lab, fn in [("Mean","AVERAGE"),("Median","MEDIAN"),("Min","MIN"),("Max","MAX")]:
        s.cell(r, 1, lab).font = Font(name=F); s.cell(r, 4, f"={fn}({rngs(colmap['Hours'])})").number_format = "0.0"; r += 1
    s.column_dimensions["A"].width = 44
    for L in "BCD": s.column_dimensions[L].width = 10
    for row in s.iter_rows(min_row=3):
        for c in row:
            if c.font is None or not c.font.bold: c.font = Font(name=F)

    notes = wb.create_sheet("Notes")
    lines = [
     ("Synthetic time entries v2: corporate/transactional practice", True),
     ("All clients, people, matters, and entries are fictional. Generated by generate_synthetic_v2.py (seed 210); rerun to regenerate or scale.", False),
     ("", False),
     ("HOW TO USE", True),
     ("Blue-header columns in Time Entries are what the timekeeper submitted. Gold 'KEY:' columns are the answer key. Drop KEY columns before feeding a model.", False),
     ("Matter Master and TK Matter History are reference tables a real reviewer would have. Models may use them.", False),
     ("Train/test: use KEY: Split. TEST rows use narrative phrasings never seen in TRAIN and client misspellings never seen in TRAIN or in Matter Master. Do not random-split.", False),
     ("", False),
     ("EXPECTED ACTION (primary label)", True),
     ("AUTO_MATCH: one matter is supported by the entered fields plus narrative. Score against KEY: True CMID.", False),
     ("FLAG_MISMATCH: entered CMID is closed or contradicts the narrative. Correct output is a flag plus the suggested matter (True CMID).", False),
     ("NEEDS_REVIEW: several matters remain plausible. Correct output is routing to a human. True CMID is kept for reference only; do not score a guess as correct.", False),
     ("SPLIT_ENTRY: narrative covers two matters. True CMID lists both ('A | B'). Correct output is a split flag.", False),
     ("Labels come from one rule-based oracle that uses only entered fields, matter-specific terms, task type vs. matter type, and TK history. It never breaks ties with hours.", False),
     ("", False),
     ("OTHER KEY COLUMNS", True),
     ("Input Type: which fields were filled (FULL, CMID_ONLY incl. bad formats like 6103-001 / 6103.1 / 6103001, CLIENT_ONLY, CODENAME, NONE, WRONG_CMID, INVALID_CMID).", False),
     ("Hard Case: GENUINELY_AMBIGUOUS, OFF_HISTORY (TK billing a matter not in their 90-day history), MULTI_MATTER, SAME_CLIENT_WRONG_MATTER, CLOSED_MATTER (old CMID typed from habit), INVALID_CMID.", False),
     ("Client Surface: CANONICAL, LISTED_ALIAS (in Matter Master), UNLISTED_VARIANT (generated misspelling/abbreviation not in Matter Master).", False),
     ("Narrative Quality: SPECIFIC, VAGUE, BLOCK_BILLED. Provided for convenience only; Stage 2 ground truth still comes from fee examiner reports.", False),
     ("Task IDs / Template IDs: which task template and phrasing produced the narrative. Use for grouped splits and error analysis.", False),
     ("", False),
     ("KNOWN LIMITATIONS (be upfront about these)", True),
     ("Templates and noise are written by us, so real entries will be messier and more varied than any paraphrase set.", False),
     ("Every FLAG_MISMATCH is detectable from the narrative by design; real wrong-CMID entries with generic narratives are undetectable and were excluded.", False),
     ("Tabular input only. OCR/email/handwritten-log text (multiple entries per line, '1:30' hours, relative dates, OCR character errors) is a planned v3.", False),
     ("Single practice group (corporate/transactional); bankruptcy fee applications used in Stage 2 are a different domain.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = notes.cell(i, 1, t); c.font = Font(name=F, bold=b, size=12 if i == 1 else 10); c.alignment = Alignment(wrap_text=True)
    notes.column_dimensions["A"].width = 130
    wb.move_sheet("Notes", offset=-4)
    wb.save(OUT); print("saved", OUT, len(TE))
